"""Shared pytest fixtures. Generates tiny sample files at runtime (no binaries committed)."""
from pathlib import Path

import pymupdf
import pytest

# No word starts with a capital I: in the default sans-serif face 'I' and 'l'
# are the same glyph, so an OCR miss there measures the font, not the pipeline.
SCANNED_LINES = (
    "Cerere de concediu de odihna",
    "Angajat: Maria Popescu",
    "Aprobat de managerul de departament",
)


@pytest.fixture
def scanned_pdf(tmp_path: Path) -> Path:
    """A 1-page PDF whose only content is a picture of text.

    Built by rendering real text, rasterizing it, then embedding the bitmap in
    a fresh text-free page — the same thing a scanner produces. Tests assert
    ``page_class == 'scanned'`` before asking OCR to read it, so the fixture
    cannot quietly stop being a scan.
    """
    source = pymupdf.open()
    page = source.new_page()
    for index, line in enumerate(SCANNED_LINES):
        page.insert_text((72, 150 + index * 40), line, fontsize=16)
    pixmap = page.get_pixmap(dpi=200)
    source.close()

    scanned = pymupdf.open()
    target = scanned.new_page(width=pixmap.width * 0.72, height=pixmap.height * 0.72)
    target.insert_image(target.rect, pixmap=pixmap)
    out = tmp_path / "scanned_test.pdf"
    scanned.save(out)
    scanned.close()
    return out


@pytest.fixture
def scanned_image(tmp_path: Path) -> Path:
    """A PNG holding the same picture of text as ``scanned_pdf``."""
    source = pymupdf.open()
    page = source.new_page()
    for index, line in enumerate(SCANNED_LINES):
        page.insert_text((72, 150 + index * 40), line, fontsize=16)
    out = tmp_path / "scanned_test.png"
    page.get_pixmap(dpi=200).save(out)
    source.close()
    return out


@pytest.fixture
def layout_complex_pdf(tmp_path: Path) -> Path:
    """A 1-page PDF with real text, many ruling lines, and one small picture.

    Mimics a diagram-heavy slide: legitimate native text (classification must
    not fall back to 'garbled'/'scanned'), more ruling lines than
    ``RULING_LINE_COUNT_HIGH`` (classifies as 'layout-complex'), and a small
    embedded picture of text — one covering under ``IMAGE_AREA_RATIO_HIGH``
    of the page, so it stays 'layout-complex' rather than becoming 'mixed'.
    """
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Layout complex slide with diagram lines and an image.",
                      fontsize=14)

    shape = page.new_shape()
    for index in range(35):
        shape.draw_line((50, 100 + index * 5), (500, 100 + index * 5))
    shape.finish(color=(0, 0, 0), width=0.5)
    shape.commit()

    picture = pymupdf.open()
    picture_page = picture.new_page()
    for index, line in enumerate(SCANNED_LINES):
        picture_page.insert_text((10, 30 + index * 40), line, fontsize=16)
    page.insert_image(pymupdf.Rect(400, 400, 470, 430), pixmap=picture_page.get_pixmap(dpi=150))
    picture.close()

    out = tmp_path / "layout_complex_test.pdf"
    doc.save(out)
    doc.close()
    return out


@pytest.fixture
def layout_complex_pdf_no_image(tmp_path: Path) -> Path:
    """Same as ``layout_complex_pdf`` but with no embedded picture at all."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Layout complex slide with diagram lines, no image.",
                      fontsize=14)

    shape = page.new_shape()
    for index in range(35):
        shape.draw_line((50, 100 + index * 5), (500, 100 + index * 5))
    shape.finish(color=(0, 0, 0), width=0.5)
    shape.commit()

    out = tmp_path / "layout_complex_no_image_test.pdf"
    doc.save(out)
    doc.close()
    return out


@pytest.fixture(autouse=True)
def reset_ocr_state():
    """Keep each test's OCR configuration and cached probe to itself."""
    from extractor.ocr import config

    config.reset()
    yield
    config.reset()


@pytest.fixture
def sample_pdf(tmp_path: Path) -> Path:
    """A 2-page text PDF with predictable content."""
    doc = pymupdf.open()
    page1 = doc.new_page()
    page1.insert_text((72, 72), "Procedura de concedii aprobata de manager.")
    page2 = doc.new_page()
    page2.insert_text((72, 72), "A doua pagina cu text simplu.")
    out = tmp_path / "procedura_test.pdf"
    doc.save(out)
    doc.close()
    return out


@pytest.fixture
def table_pdf(tmp_path: Path) -> Path:
    """A 1-page PDF with a paragraph plus a ruled 2-col, 3-row table."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Procedura de concedii. Managerul aproba cererea.")
    x0, y0, w, h, rows, cols = 72, 120, 200, 60, 3, 2
    cw, rh = w / cols, h / rows
    for r in range(rows + 1):
        page.draw_line((x0, y0 + r * rh), (x0 + w, y0 + r * rh))
    for c in range(cols + 1):
        page.draw_line((x0 + c * cw, y0), (x0 + c * cw, y0 + h))
    cells = [["Rol", "Actiune"], ["Manager", "Aproba"], ["HR", "Inregistreaza"]]
    for r in range(rows):
        for c in range(cols):
            page.insert_text((x0 + c * cw + 5, y0 + r * rh + 15), cells[r][c])
    out = tmp_path / "table_test.pdf"
    doc.save(out)
    doc.close()
    return out


@pytest.fixture
def sample_csv(tmp_path: Path) -> Path:
    """A 2-column, 3-row CSV with a header."""
    out = tmp_path / "date_test.csv"
    out.write_text(
        "Rol,Actiune\nManager,Aproba\nHR,Inregistreaza\n",
        encoding="utf-8",
    )
    return out


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    """A Word doc: a paragraph, then a 2x2 table, then a paragraph."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("Procedura de concedii aprobata de manager.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Rol"
    table.cell(0, 1).text = "Actiune"
    table.cell(1, 0).text = "Manager"
    table.cell(1, 1).text = "Aproba"
    doc.add_paragraph("A doua sectiune cu text simplu.")
    out = tmp_path / "raport_test.docx"
    doc.save(out)
    return out


@pytest.fixture
def sample_xlsx(tmp_path: Path) -> Path:
    """A workbook with one populated sheet and one empty sheet."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Date"
    ws.append(["Rol", "Actiune"])
    ws.append(["Manager", "Aproba"])
    ws.append(["HR", "Inregistreaza"])
    wb.create_sheet("Goala")  # left empty on purpose
    out = tmp_path / "registru_test.xlsx"
    wb.save(out)
    return out
