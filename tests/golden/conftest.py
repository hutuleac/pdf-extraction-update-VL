"""Golden corpus fixtures — deterministically generated edge-case documents.

Following the project convention, all samples are constructed at runtime with
PyMuPDF / python-docx / openpyxl so no binary fixtures are committed.
"""
from pathlib import Path

import pymupdf
import pytest

# ---------------------------------------------------------------------------
# PDF fixtures
# ---------------------------------------------------------------------------

@pytest.fixture
def single_column_pdf(tmp_path: Path) -> Path:
    """A 3-page single-column text PDF."""
    doc = pymupdf.open()
    for i in range(1, 4):
        page = doc.new_page()
        # Place text well below the top band (12% of 842pt = ~101pt)
        page.insert_text((72, 150), f"Pagina {i}: Continut text simplu pe o singura coloana.")
        page.insert_text((72, 180), f"Linia a doua pe pagina {i} cu detalii unice.")
        page.insert_text((72, 210), f"Informatii specifice paginii numarul {i}.")
    out = tmp_path / "single_column.pdf"
    doc.save(out)
    doc.close()
    return out


@pytest.fixture
def two_column_pdf(tmp_path: Path) -> Path:
    """A 3-page two-column PDF with distinct left/right content."""
    doc = pymupdf.open()
    for i in range(1, 4):
        page = doc.new_page()
        # Left column (x=72)
        page.insert_text((72, 72), f"Stanga pagina {i} linia 1.")
        page.insert_text((72, 90), f"Stanga pagina {i} linia 2.")
        page.insert_text((72, 108), f"Stanga pagina {i} linia 3.")
        # Right column (x=350)
        page.insert_text((350, 72), f"Dreapta pagina {i} linia 1.")
        page.insert_text((350, 90), f"Dreapta pagina {i} linia 2.")
        page.insert_text((350, 108), f"Dreapta pagina {i} linia 3.")
    out = tmp_path / "two_column.pdf"
    doc.save(out)
    doc.close()
    return out


@pytest.fixture
def image_only_pdf(tmp_path: Path) -> Path:
    """A 2-page PDF with only images (no text) — simulates scanned pages."""
    doc = pymupdf.open()
    for _ in range(2):
        page = doc.new_page()
        # Insert a filled rectangle as a stand-in for an image
        rect = pymupdf.Rect(50, 50, 500, 700)
        page.draw_rect(rect, fill=(0.9, 0.9, 0.9))
        # Insert a tiny pixmap as an actual image
        pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 100, 100), 1)
        pix.clear_with(200)
        page.insert_image(pymupdf.Rect(100, 100, 400, 600), pixmap=pix)
    out = tmp_path / "image_only.pdf"
    doc.save(out)
    doc.close()
    return out


@pytest.fixture
def mixed_content_pdf(tmp_path: Path) -> Path:
    """A 2-page PDF with text AND large images on each page."""
    doc = pymupdf.open()
    for i in range(1, 3):
        page = doc.new_page()
        page.insert_text((72, 72), f"Text pe pagina {i} cu continut mixt.")
        page.insert_text((72, 90), "Mai mult text sub imagine.")
        # Large image covering most of the page
        pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 200), 1)
        pix.clear_with(150)
        page.insert_image(pymupdf.Rect(50, 150, 500, 700), pixmap=pix)
    out = tmp_path / "mixed_content.pdf"
    doc.save(out)
    doc.close()
    return out


@pytest.fixture
def repeated_header_pdf(tmp_path: Path) -> Path:
    """A 6-page PDF with 'Confidential | Page N' footer on every page."""
    doc = pymupdf.open()
    for i in range(1, 7):
        page = doc.new_page()
        # Main content in the middle of the page (below top band)
        page.insert_text((72, 200), f"Continut pagina {i}.")
        page.insert_text((72, 230), f"Detalii importante pe pagina {i}.")
        page.insert_text((72, 260), f"Sectiune unica pentru pagina {i} cu date specifice.")
        # Footer at bottom of page (in bottom band)
        page.insert_text((72, 780), f"Confidential | Page {i}")
    out = tmp_path / "repeated_header.pdf"
    doc.save(out)
    doc.close()
    return out


@pytest.fixture
def unicode_pdf(tmp_path: Path) -> Path:
    """A 2-page PDF with Romanian diacritics and special Unicode chars."""
    doc = pymupdf.open()
    page1 = doc.new_page()
    page1.insert_text((72, 72), "Înregistrarea cererii de concediu.")
    page1.insert_text((72, 90), "Managerul aprobă solicitarea.")
    page1.insert_text((72, 108), "Departamentul HR verifică disponibilitatea.")
    page2 = doc.new_page()
    page2.insert_text((72, 72), "Notă: Toți angajații trebuie să completeze formularul.")
    page2.insert_text((72, 90), "Valabil începând cu 1 ianuarie 2026.")
    out = tmp_path / "unicode_test.pdf"
    doc.save(out)
    doc.close()
    return out


# ---------------------------------------------------------------------------
# DOCX fixtures
# ---------------------------------------------------------------------------

@pytest.fixture
def nested_list_docx(tmp_path: Path) -> Path:
    """A DOCX with nested lists between paragraphs."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("Introducere document.")
    doc.add_paragraph("Punct 1", style="List Bullet")
    doc.add_paragraph("Punct 2", style="List Bullet")
    doc.add_paragraph("Sub-punct 2a", style="List Bullet 2")
    doc.add_paragraph("Punct 3", style="List Bullet")
    doc.add_paragraph("Concluzie document.")
    out = tmp_path / "nested_list.docx"
    doc.save(out)
    return out


@pytest.fixture
def table_between_paragraphs_docx(tmp_path: Path) -> Path:
    """A DOCX with paragraphs, a table, then more paragraphs."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("Text inainte de tabel.")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Nume"
    table.cell(0, 1).text = "Departament"
    table.cell(1, 0).text = "Ion Popescu"
    table.cell(1, 1).text = "IT"
    table.cell(2, 0).text = "Maria Ionescu"
    table.cell(2, 1).text = "HR"
    doc.add_paragraph("Text dupa tabel.")
    doc.add_paragraph("Paragraf final.")
    out = tmp_path / "table_between.docx"
    doc.save(out)
    return out


# ---------------------------------------------------------------------------
# XLSX fixtures
# ---------------------------------------------------------------------------

@pytest.fixture
def merged_cell_xlsx(tmp_path: Path) -> Path:
    """An XLSX with merged cells."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Merged"
    ws.merge_cells("A1:C1")
    ws["A1"] = "Header Combinat"
    ws["A2"] = "Val1"
    ws["B2"] = "Val2"
    ws["C2"] = "Val3"
    ws["A3"] = "Data1"
    ws["B3"] = "Data2"
    ws["C3"] = "Data3"
    out = tmp_path / "merged_cells.xlsx"
    wb.save(out)
    return out


@pytest.fixture
def hidden_sheet_xlsx(tmp_path: Path) -> Path:
    """An XLSX with one visible sheet and one hidden sheet."""
    from openpyxl import Workbook

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Visible"
    ws1.append(["Col1", "Col2"])
    ws1.append(["A", "B"])
    ws2 = wb.create_sheet("Hidden")
    ws2.append(["Secret", "Data"])
    ws2.append(["X", "Y"])
    ws2.sheet_state = "hidden"
    out = tmp_path / "hidden_sheet.xlsx"
    wb.save(out)
    return out


@pytest.fixture
def dates_and_currency_xlsx(tmp_path: Path) -> Path:
    """An XLSX with dates and currency values."""
    from datetime import date

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Financial"
    ws.append(["Date", "Amount", "Currency"])
    ws.append([date(2026, 1, 15), 1500.50, "EUR"])
    ws.append([date(2026, 2, 28), 2300.75, "RON"])
    ws.append([date(2026, 3, 10), 890.00, "USD"])
    out = tmp_path / "dates_currency.xlsx"
    wb.save(out)
    return out
