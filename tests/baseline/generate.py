"""Generate the committed pre-OCR output snapshot.

The snapshot is what proves OCR changed nothing for documents that never
needed it. It is built from deterministic fixtures rather than from ``input/``,
because that folder holds customer material that must not reach the remote.

Only documents with no scanned or image-dominant pages are included: those are
the ones whose output must stay byte-identical forever. Scanned and mixed pages
are covered by the OCR tests instead, since their output is meant to change.

    python tests/baseline/generate.py          # rewrite the snapshot
"""
from __future__ import annotations

import json
import sys
import tempfile
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from extractor.dispatcher import extract_document

SNAPSHOT_DIR = Path(__file__).parent


def build_single_column_pdf(directory: Path) -> Path:
    """A 3-page single-column text PDF."""
    import pymupdf

    doc = pymupdf.open()
    for i in range(1, 4):
        page = doc.new_page()
        page.insert_text((72, 150), f"Pagina {i}: Continut text simplu pe o singura coloana.")
        page.insert_text((72, 180), f"Linia a doua pe pagina {i} cu detalii unice.")
        page.insert_text((72, 210), f"Informatii specifice paginii numarul {i}.")
    out = directory / "single_column.pdf"
    doc.save(out)
    doc.close()
    return out


def build_two_column_pdf(directory: Path) -> Path:
    """A 3-page two-column PDF with distinct left/right content."""
    import pymupdf

    doc = pymupdf.open()
    for i in range(1, 4):
        page = doc.new_page()
        page.insert_text((72, 72), f"Stanga pagina {i} linia 1.")
        page.insert_text((72, 90), f"Stanga pagina {i} linia 2.")
        page.insert_text((72, 108), f"Stanga pagina {i} linia 3.")
        page.insert_text((350, 72), f"Dreapta pagina {i} linia 1.")
        page.insert_text((350, 90), f"Dreapta pagina {i} linia 2.")
        page.insert_text((350, 108), f"Dreapta pagina {i} linia 3.")
    out = directory / "two_column.pdf"
    doc.save(out)
    doc.close()
    return out


def build_repeated_header_pdf(directory: Path) -> Path:
    """A 6-page PDF with a 'Confidential | Page N' footer on every page."""
    import pymupdf

    doc = pymupdf.open()
    for i in range(1, 7):
        page = doc.new_page()
        page.insert_text((72, 200), f"Continut pagina {i}.")
        page.insert_text((72, 230), f"Detalii importante pe pagina {i}.")
        page.insert_text((72, 260), f"Sectiune unica pentru pagina {i} cu date specifice.")
        page.insert_text((72, 780), f"Confidential | Page {i}")
    out = directory / "repeated_header.pdf"
    doc.save(out)
    doc.close()
    return out


def build_unicode_pdf(directory: Path) -> Path:
    """A 2-page PDF with Romanian diacritics and special Unicode characters."""
    import pymupdf

    doc = pymupdf.open()
    page1 = doc.new_page()
    page1.insert_text((72, 72), "Înregistrarea cererii de concediu.")
    page1.insert_text((72, 90), "Managerul aprobă solicitarea.")
    page1.insert_text((72, 108), "Departamentul HR verifică disponibilitatea.")
    page2 = doc.new_page()
    page2.insert_text((72, 72), "Notă: Toți angajații trebuie să completeze formularul.")
    page2.insert_text((72, 90), "Valabil începând cu 1 ianuarie 2026.")
    out = directory / "unicode_test.pdf"
    doc.save(out)
    doc.close()
    return out


def build_nested_list_docx(directory: Path) -> Path:
    """A DOCX with nested lists between paragraphs."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("Introducere document.")
    doc.add_paragraph("Punct 1", style="List Bullet")
    doc.add_paragraph("Punct 2", style="List Bullet")
    doc.add_paragraph("Sub-punct 2a", style="List Bullet 2")
    doc.add_paragraph("Punct 3", style="List Bullet")
    doc.add_paragraph("Concluzie document.")
    out = directory / "nested_list.docx"
    doc.save(out)
    return out


def build_table_between_docx(directory: Path) -> Path:
    """A DOCX with paragraphs, a table, then more paragraphs."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("Text inainte de tabel.")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Nume"
    table.cell(0, 1).text = "Departament"
    table.cell(1, 0).text = "Ion Popescu"
    table.cell(1, 1).text = "IT"
    table.cell(2, 0).text = "Maria Ionescu"
    table.cell(2, 1).text = "HR"
    doc.add_paragraph("Text dupa tabel.")
    doc.add_paragraph("Paragraf final.")
    out = directory / "table_between.docx"
    doc.save(out)
    return out


def build_merged_cell_xlsx(directory: Path) -> Path:
    """An XLSX with merged cells."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Merged"
    ws.merge_cells("A1:C1")
    ws["A1"] = "Header Combinat"
    ws["A2"], ws["B2"], ws["C2"] = "Val1", "Val2", "Val3"
    ws["A3"], ws["B3"], ws["C3"] = "Data1", "Data2", "Data3"
    out = directory / "merged_cells.xlsx"
    wb.save(out)
    return out


def build_hidden_sheet_xlsx(directory: Path) -> Path:
    """An XLSX with one visible sheet and one hidden sheet."""
    from openpyxl import Workbook

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Visible"
    ws1.append(["Col1", "Col2"])
    ws1.append(["A", "B"])
    ws2 = wb.create_sheet("Hidden")
    ws2.append(["Secret", "Data"])
    ws2.append(["X", "Y"])
    ws2.sheet_state = "hidden"
    out = directory / "hidden_sheet.xlsx"
    wb.save(out)
    return out


def build_dates_currency_xlsx(directory: Path) -> Path:
    """An XLSX with dates and currency values."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Financial"
    ws.append(["Date", "Amount", "Currency"])
    ws.append([date(2026, 1, 15), 1500.50, "EUR"])
    ws.append([date(2026, 2, 28), 2300.75, "RON"])
    ws.append([date(2026, 3, 10), 890.00, "USD"])
    out = directory / "dates_currency.xlsx"
    wb.save(out)
    return out


BUILDERS = (
    build_single_column_pdf,
    build_two_column_pdf,
    build_repeated_header_pdf,
    build_unicode_pdf,
    build_nested_list_docx,
    build_table_between_docx,
    build_merged_cell_xlsx,
    build_hidden_sheet_xlsx,
    build_dates_currency_xlsx,
)


def build_snapshot() -> dict[str, dict]:
    """Extract every baseline fixture and return {fixture name: model}."""
    snapshot: dict[str, dict] = {}
    with tempfile.TemporaryDirectory() as raw_dir:
        directory = Path(raw_dir)
        for builder in BUILDERS:
            path = builder(directory)
            snapshot[path.name] = extract_document(path)
    return snapshot


def write_snapshot() -> Path:
    """Write the snapshot to ``tests/baseline/snapshot.json`` and return the path."""
    out = SNAPSHOT_DIR / "snapshot.json"
    payload = json.dumps(build_snapshot(), ensure_ascii=False, indent=2)
    out.write_text(payload + "\n", encoding="utf-8")
    return out


if __name__ == "__main__":
    written = write_snapshot()
    print(f"Wrote {written}")
