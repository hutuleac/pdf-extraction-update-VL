from openpyxl import Workbook

from extractor.xlsx_reader import extract_xlsx


def test_extract_xlsx_shape_skips_empty(sample_xlsx):
    model = extract_xlsx(sample_xlsx)
    assert model["document"]["source_type"] == "xlsx"
    assert model["document"]["filename"] == "registru_test.xlsx"
    # "Goala" is empty and skipped -> one sheet unit remains.
    assert model["document"]["pages"] == 1
    unit = model["pages"][0]
    assert unit["unit"] == 1
    assert unit["unit_type"] == "sheet"


def test_extract_xlsx_values(sample_xlsx):
    blocks = extract_xlsx(sample_xlsx)["pages"][0]["content"]
    tables = [b for b in blocks if b["type"] == "table"]
    assert tables[0]["content"] == [
        ["Rol", "Actiune"],
        ["Manager", "Aproba"],
        ["HR", "Inregistreaza"],
    ]


def test_extract_xlsx_reads_computed_values_not_formulas(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 2
    ws["B1"] = 3
    ws["C1"] = "=A1+B1"
    # openpyxl stores the formula string; there is no cached value until a
    # spreadsheet app recomputes it. data_only therefore yields None -> "".
    out = tmp_path / "formula.xlsx"
    wb.save(out)

    blocks = extract_xlsx(out)["pages"][0]["content"]
    rows = blocks[0]["content"]
    assert rows[0][2] == ""  # not the literal "=A1+B1"
