"""PowerPoint (.pptx) extraction via python-pptx.

One `slide` unit per slide. Text frames in shape order become text blocks,
graphic_frame tables become table blocks, speaker notes are a separate text
block (included by default). Pictures are read by OCR through the same
`apply.ocr_images` glue the PDF and image readers use — a screenshot-only
slide is the normal case in real decks, not an edge case.

Group shapes are traversed recursively: grouping is standard practice, and the
text inside a group is ordinary slide text.
"""
import logging
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from extractor.model import make_document, make_table_block, make_text_block, make_unit

logger = logging.getLogger(__name__)


def _extract_table_from_shape(shape) -> list[list[str]] | None:
    """Extract a table from a shape's graphic_frame, if it has one."""
    if not shape.has_table:
        return None
    table = shape.table
    rows: list[list[str]] = []
    for row in table.rows:
        cells = [cell.text for cell in row.cells]
        rows.append(cells)
    return rows if rows else None


def _chart_to_rows(chart) -> list[list[str]] | None:
    """Flatten a native chart's data into table rows: categories down, series across.

    The values are the deck author's own numbers, so this is the only lossless
    read of a chart — OCR sees a rendered bar, the VLM guesses at it.
    """
    try:
        plot = chart.plots[0]
        categories = [str(c) if c is not None else "" for c in plot.categories]
        series = list(plot.series)
    except (IndexError, AttributeError, ValueError):
        return None
    if not series:
        return None
    header = [""] + [s.name or "" for s in series]
    columns = [list(s.values) for s in series]
    n_rows = max(len(categories), *(len(col) for col in columns))
    rows = [header]
    for i in range(n_rows):
        label = categories[i] if i < len(categories) else ""
        cells = [_fmt(col[i]) if i < len(col) else "" for col in columns]
        rows.append([label, *cells])
    return rows


def _fmt(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _iter_shapes(shapes):
    """Yield every shape, descending into groups.

    A grouped object's children are not in `slide.shapes`, so without this
    every word inside a group is silently dropped.
    """
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _extract_slide_text(slide) -> tuple[list[dict], list[bytes]]:
    """Extract blocks from a slide in shape order, plus its picture blobs."""
    blocks: list[dict] = []
    pictures: list[bytes] = []

    for shape in _iter_shapes(slide.shapes):
        # Table shapes
        if shape.has_table:
            rows = _extract_table_from_shape(shape)
            if rows:
                blocks.append(make_table_block(rows))
        elif getattr(shape, "has_chart", False):
            chart = shape.chart
            if chart.has_title and chart.chart_title.text_frame.text.strip():
                blocks.append(make_text_block(chart.chart_title.text_frame.text))
            rows = _chart_to_rows(chart)
            if rows:
                blocks.append(make_table_block(rows, source="chart"))
        # Text frame shapes (text boxes, titles, etc.)
        elif shape.has_text_frame:
            text = shape.text_frame.text
            text_block = make_text_block(text)
            if text_block:
                blocks.append(text_block)
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            try:
                pictures.append(shape.image.blob)
            except Exception as exc:  # noqa: BLE001 - one bad picture is not a failed slide
                logger.debug("Could not read a slide picture: %s", exc)

    return blocks, pictures


def _ocr_pictures(pictures: list[bytes], slide_number: int) -> tuple[dict | None, list[dict]]:
    """OCR a slide's pictures via the shared glue, or explain why not.

    Mirrors the PDF reader: unavailability is a structured warning, never a
    crash, and a slide with no pictures asks for nothing.
    """
    from extractor.ocr import registry
    from extractor.ocr.apply import ocr_images, unavailable_warning
    from extractor.ocr.config import get_config
    from extractor.ocr.raster import image_bytes_to_array

    if not pictures:
        return None, []
    if not registry.is_available():
        return None, [unavailable_warning(slide_number)]

    max_pixels = get_config().max_pixels
    images = []
    for blob in pictures:
        try:
            images.append(image_bytes_to_array(blob, max_pixels=max_pixels))
        except Exception as exc:  # noqa: BLE001 - isolate one picture
            logger.warning("Could not decode a picture on slide %d: %s", slide_number, exc)
    if not images:
        return None, [{"code": "OCR_FAILED", "page": slide_number,
                       "detail": "no picture on this slide could be decoded"}]

    return ocr_images(images, slide_number)


def _extract_notes(slide) -> dict | None:
    """Extract speaker notes as a text block, or None if empty."""
    if not slide.has_notes_slide:
        return None
    notes_slide = slide.notes_slide
    notes_text = notes_slide.notes_text_frame.text if notes_slide.notes_text_frame else ""
    if not notes_text or not notes_text.strip():
        return None
    return make_text_block(notes_text)


def extract_pptx(path: Path | str) -> dict:
    """Extract a PowerPoint file into the internal model (one unit per slide)."""
    path = Path(path)
    prs = Presentation(path)

    units: list[dict] = []
    warnings: list[dict] = []
    total_images = 0
    for slide_number, slide in enumerate(prs.slides, start=1):
        blocks, pictures = _extract_slide_text(slide)
        total_images += len(pictures)

        ocr_block, ocr_warnings = _ocr_pictures(pictures, slide_number)
        if ocr_block:
            blocks.append(ocr_block)
        warnings.extend(ocr_warnings)

        # Append speaker notes as a separate text block
        notes_block = _extract_notes(slide)
        if notes_block:
            blocks.append(notes_block)

        units.append(make_unit(
            slide_number, "slide", blocks, image_count=len(pictures),
        ))

    return make_document(
        path.name, "pptx", len(units), units,
        image_count=total_images, warnings=warnings or None,
    )
