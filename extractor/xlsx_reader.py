"""Excel (.xlsx) extraction via openpyxl.

Each worksheet becomes one `sheet` unit holding a single table block of cell
values. `data_only=True` reads computed values rather than formula strings,
which is what phase-2 embeddings want. Fully empty sheets are skipped.
"""
import logging
from pathlib import Path

from openpyxl import load_workbook

from extractor.limits import (
    MAX_WORKBOOK_CELLS,
    MAX_WORKBOOK_ROWS,
    WorkbookTooLargeError,
)
from extractor.model import make_document, make_table_block, make_unit

logger = logging.getLogger(__name__)


def _preflight(workbook) -> tuple[int, int]:
    """Return the workbook's declared (rows, cells) without reading any of it.

    openpyxl exposes each sheet's dimension record even in read-only mode, so
    an oversized workbook can be refused before a single row is materialized.
    Returns zeros for the dimensions it cannot determine — a sheet with no
    dimension record reports ``max_row`` as None, and a hostile file can simply
    understate it, which is why the streaming count in ``_sheet_rows`` still
    has to hold the line.
    """
    rows = 0
    cells = 0
    for worksheet in workbook.worksheets:
        sheet_rows = worksheet.max_row or 0
        sheet_cols = worksheet.max_column or 0
        rows += sheet_rows
        cells += sheet_rows * sheet_cols
    return rows, cells


def _sheet_rows(worksheet, path: Path, counted: list[int]) -> list[list[str]]:
    """Return a worksheet's used range as rows of cell strings (None -> '').

    ``counted`` carries the running (rows, cells) total across the workbook, so
    the caps bound the whole document rather than each sheet in isolation. The
    check happens as rows arrive: the point is to stop accumulating, so
    counting first and checking afterwards would defeat it.
    """
    rows: list[list[str]] = []
    for row in worksheet.iter_rows(values_only=True):
        counted[0] += 1
        counted[1] += len(row)
        if counted[0] > MAX_WORKBOOK_ROWS or counted[1] > MAX_WORKBOOK_CELLS:
            raise WorkbookTooLargeError(path, counted[0], counted[1])
        rows.append(["" if cell is None else str(cell) for cell in row])
    return rows


def _is_empty(rows: list[list[str]]) -> bool:
    """True if the sheet has no non-blank cell."""
    return not any(any(cell for cell in row) for row in rows)


def extract_xlsx(path: Path | str) -> dict:
    """Extract an Excel workbook into the internal model (one unit per sheet)."""
    path = Path(path)
    workbook = load_workbook(path, read_only=True, data_only=True)

    units: list[dict] = []
    counted = [0, 0]
    try:
        declared_rows, declared_cells = _preflight(workbook)
        if declared_rows > MAX_WORKBOOK_ROWS or declared_cells > MAX_WORKBOOK_CELLS:
            raise WorkbookTooLargeError(path, declared_rows, declared_cells)

        for worksheet in workbook.worksheets:
            rows = _sheet_rows(worksheet, path, counted)
            if _is_empty(rows):
                continue
            blocks = [make_table_block(rows)]
            units.append(make_unit(len(units) + 1, "sheet", blocks, title=worksheet.title))
    finally:
        workbook.close()

    # Report the real extension: a macro workbook is not an .xlsx, and a
    # downstream consumer has no other way to tell the two apart.
    source_type = path.suffix.lstrip(".").lower() or "xlsx"
    return make_document(path.name, source_type, len(units), units)
